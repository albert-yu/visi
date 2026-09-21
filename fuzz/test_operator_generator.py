import re

import openpyxl
import pytest
from fuzz_excel import ExcelFuzzGenerator

OPERATORS = (
    *ExcelFuzzGenerator.SYMBOLIC_OPERATORS,
    *ExcelFuzzGenerator.PREFIX_OPERATORS,
)


def test_operator_inventory_matches_parser():
    ExcelFuzzGenerator._check_symbolic_operator_generators()


@pytest.mark.parametrize("change", ["missing", "unknown"])
def test_operator_inventory_rejects_drift(monkeypatch, change):
    ops = ExcelFuzzGenerator.SYMBOLIC_OPERATORS
    ops = ops[1:] if change == "missing" else (*ops, "Unknown")
    monkeypatch.setattr(ExcelFuzzGenerator, "SYMBOLIC_OPERATORS", ops)
    with pytest.raises(AssertionError, match="coverage mismatch"):
        ExcelFuzzGenerator._check_symbolic_operator_generators()


@pytest.mark.parametrize("op", OPERATORS)
def test_operator_cases_are_random_nested_and_reproducible(op):
    formulas = set()
    for seed in range(24):
        gen = ExcelFuzzGenerator(seed=seed)
        formula = gen.generate_symbolic_operator_formula(op, 27)
        formulas.add(formula)
        assert formula == ExcelFuzzGenerator(
            seed=seed
        ).generate_symbolic_operator_formula(op, 27)
        assert formula.startswith("=")
        assert len(formula) < 10000
        depth = 0
        max_depth = 0
        for char in formula:
            if char == "(":
                depth += 1
                max_depth = max(max_depth, depth)
            elif char == ")":
                depth -= 1
                assert depth >= 0
        assert depth == 0
        assert max_depth >= 2
        for col, row in re.findall(r"\$?([A-Z]{1,3})\$?(\d+)", formula):
            assert col in ("AA", "AB", "AC")
            assert 1 <= int(row) <= 3
    assert len(formulas) >= 8


@pytest.mark.parametrize("op,symbol", ExcelFuzzGenerator.BINARY_OPERATORS.items())
def test_binary_emitter_keeps_requested_operator(op, symbol):
    gen = ExcelFuzzGenerator(seed=1)
    formula = gen._generate_operator_expr(
        op, lambda: "(1 + 2)", lambda: "A1", lambda: "A1:A2"
    )
    assert formula.startswith(f"((1 + 2) {symbol} ")


@pytest.mark.parametrize(
    "op,pattern",
    [
        ("UnaryPlus", r"^\(\+\("),
        ("UnaryMinus", r"^\(-\("),
        ("Percent", r"%\)$"),
        ("ImplicitIntersection", r"^\(@\("),
        ("Spill", r"_xlfn\.ANCHORARRAY\(A1\)"),
        ("Intersect", r" A1:A2\)\)$"),
        ("Union", r","),
    ],
)
def test_nonbinary_emitter_keeps_requested_operator(op, pattern):
    gen = ExcelFuzzGenerator(seed=1)
    formula = gen._generate_operator_expr(
        op, lambda: "(1 + 2)", lambda: "A1", lambda: "A1:A2"
    )
    assert re.search(pattern, formula)
    assert "IFERROR" not in formula


def test_unknown_operator_is_not_silently_replaced():
    with pytest.raises(AssertionError, match="no generator wired up"):
        ExcelFuzzGenerator(seed=1).generate_symbolic_operator_formula("Unknown", 1)


def test_general_formula_trees_use_shared_operator_emitter(monkeypatch):
    gen = ExcelFuzzGenerator(seed=728)
    seen = set()
    emit = gen._generate_operator_expr

    def record(op, *args):
        seen.add(op)
        return emit(op, *args)

    monkeypatch.setattr(gen, "_generate_operator_expr", record)
    for _ in range(2000):
        gen.generate_formula(10, 5, 10, 5)
    assert seen == set(OPERATORS)


@pytest.mark.parametrize("rows,cols", [(4, 2), (10, 5)])
def test_every_workbook_contains_generated_operator_cases(
    tmp_path, monkeypatch, rows, cols
):
    gen = ExcelFuzzGenerator(seed=938)
    cases = {}
    generate = gen.generate_symbolic_operator_formula

    def record(op, data_col):
        cases[op] = generate(op, data_col)
        return cases[op]

    monkeypatch.setattr(gen, "generate_symbolic_operator_formula", record)
    path = tmp_path / "operators.xlsx"
    gen.create_fuzz_workbook(path, num_rows=rows, num_cols=cols)
    assert set(cases) == set(OPERATORS)
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb["Sheet1"]
        assert [
            ws.cell(row, ws.max_column).value for row in range(1, len(OPERATORS) + 1)
        ] == list(cases.values())
        for row in range(1, 4):
            for col in range(ws.max_column - 3, ws.max_column):
                assert isinstance(ws.cell(row, col).value, int)
    finally:
        wb.close()
